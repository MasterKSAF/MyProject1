from pathlib import Path
import re

from docx import Document
from docx.shared import Pt


ROOT = Path(r"C:\Users\Misha\Documents\GitHub\MyProject1")
SOURCE_PATH = ROOT / "docs" / "stage8_exam_memo_ru.md"
OUTPUT_PATH = ROOT / "docs" / "stage8_exam_memo.docx"


def normalize_line(line: str) -> str:
    line = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1: \2", line)
    line = line.replace("`", "")
    return line


def main() -> None:
    text = SOURCE_PATH.read_text(encoding="utf-8")
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    for raw_line in text.splitlines():
        line = raw_line.rstrip()
        if line.startswith("# "):
            document.add_heading(normalize_line(line[2:].strip()), level=1)
            continue
        if line.startswith("## "):
            document.add_heading(normalize_line(line[3:].strip()), level=2)
            continue
        if line.startswith("### "):
            document.add_heading(normalize_line(line[4:].strip()), level=3)
            continue
        if not line:
            document.add_paragraph("")
            continue
        if line.startswith("- "):
            document.add_paragraph(normalize_line(line[2:].strip()), style="List Bullet")
            continue
        if line.startswith(("1. ", "2. ", "3. ", "4. ", "5. ", "6. ")):
            document.add_paragraph(normalize_line(line.strip()), style="List Number")
            continue
        if line.startswith("```"):
            continue
        document.add_paragraph(normalize_line(line.strip()))

    document.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    main()
