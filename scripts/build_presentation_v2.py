from __future__ import annotations

from pathlib import Path
import shutil

import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Pt
from pptx.util import Inches


ROOT = Path(__file__).resolve().parents[1]
SOURCE_PPTX = ROOT / "Presentation.pptx"
OUTPUT_PPTX = ROOT / "Presentation_v2.pptx"
TEMP_OUTPUT_PPTX = ROOT / "Presentation_v2_tmp_build.pptx"
CHART_PATH = ROOT / "docs" / "presentation_v2_comparison.png"


def set_shape_text(shape, text: str) -> None:
    text_frame = shape.text_frame
    if not text_frame.paragraphs:
        paragraph = text_frame.add_paragraph()
    else:
        paragraph = text_frame.paragraphs[0]

    if paragraph.runs:
        paragraph.runs[0].text = text
    else:
        paragraph.add_run().text = text

    for run in paragraph.runs[1:]:
        run.text = ""

    for extra_paragraph in text_frame.paragraphs[1:]:
        for run in extra_paragraph.runs:
            run.text = ""


def style_shape_text(
    shape,
    *,
    font_size: float | None = None,
    bold: bool | None = None,
    color: tuple[int, int, int] | None = None,
    underline: bool | None = None,
) -> None:
    if not hasattr(shape, "text_frame"):
        return
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if font_size is not None:
                run.font.size = Pt(font_size)
            if bold is not None:
                run.font.bold = bold
            if underline is not None:
                run.font.underline = underline
            if color is not None:
                run.font.color.rgb = RGBColor(*color)


def add_comparison_chart(output_path: Path) -> None:
    models = ["ResNet18", "ResNet34", "MobileNetV3"]
    accuracy = [89.74, 85.71, 49.08]
    colors = ["#D1495B", "#2E5EAA", "#7A8B99"]

    plt.figure(figsize=(7.2, 3.4), dpi=200)
    ax = plt.gca()
    bars = ax.bar(models, accuracy, color=colors, width=0.58)
    ax.set_ylim(0, 100)
    ax.set_ylabel("Accuracy, %")
    ax.set_title("Сравнение моделей на 20 эпохах", fontsize=13, pad=12)
    ax.grid(axis="y", linestyle="--", alpha=0.25)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)

    for bar, value in zip(bars, accuracy):
        ax.text(
            bar.get_x() + bar.get_width() / 2,
            value + 1.2,
            f"{value:.2f}%",
            ha="center",
            va="bottom",
            fontsize=10,
            fontweight="bold",
        )

    plt.tight_layout()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    plt.savefig(output_path, bbox_inches="tight", facecolor="white")
    plt.close()


def remove_shape(shape) -> None:
    shape._element.getparent().remove(shape._element)


def add_textbox(
    slide,
    left: int,
    top: int,
    width: int,
    height: int,
    text: str,
    *,
    font_size: float,
    bold: bool = False,
    color: tuple[int, int, int] = (241, 245, 249),
) -> None:
    textbox = slide.shapes.add_textbox(left, top, width, height)
    text_frame = textbox.text_frame
    text_frame.word_wrap = True
    paragraph = text_frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.LEFT
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*color)


def main() -> None:
    add_comparison_chart(CHART_PATH)

    prs = Presentation(str(SOURCE_PPTX))

    # Slide 1
    slide = prs.slides[0]
    set_shape_text(slide.shapes[6], "Финальная версия нейронной сети · ResNet18")

    # Slide 4
    slide = prs.slides[3]
    for idx in [6, 11, 21]:
        slide.shapes[idx].top = slide.shapes[16].top
        slide.shapes[idx].height = slide.shapes[16].height

    # Slide 6
    slide = prs.slides[5]
    set_shape_text(slide.shapes[5], "ResNet34")
    set_shape_text(slide.shapes[6], "Более глубокая residual-сеть. На тесте уступила ResNet18.")
    set_shape_text(slide.shapes[7], "Batch size 256")
    set_shape_text(slide.shapes[8], "Epochs 20")
    set_shape_text(slide.shapes[9], "Параметры 21.3M")
    set_shape_text(slide.shapes[11], "85.7%")

    set_shape_text(slide.shapes[15], "ResNet18")
    set_shape_text(slide.shapes[16], "Baseline-модель. В обновлённом сравнении показала лучший результат.")
    set_shape_text(slide.shapes[17], "Batch size 256")
    set_shape_text(slide.shapes[18], "Epochs 20")
    set_shape_text(slide.shapes[19], "Параметры 11.2M")
    set_shape_text(slide.shapes[21], "89.7%")

    set_shape_text(slide.shapes[25], "Компактная мобильная модель. Сильно уступила ResNet по Accuracy.")
    set_shape_text(slide.shapes[27], "Epochs 20")
    set_shape_text(slide.shapes[30], "49.1%")
    set_shape_text(
        slide.shapes[32],
        "Модели сравнивались при image size 160×160, batch size 256 и 20 эпохах.",
    )

    winner_bg = slide.shapes[13]
    winner_text = slide.shapes[14]
    winner_bg.left = 3319425
    winner_text.left = 3986138
    set_shape_text(winner_text, "ПОБЕДИТЕЛЬ")

    # Slide 7
    slide = prs.slides[6]
    set_shape_text(slide.shapes[11], "ResNet18")
    set_shape_text(slide.shapes[12], "89.74%")
    set_shape_text(slide.shapes[13], "11.2M")
    set_shape_text(slide.shapes[16], "ResNet34")
    set_shape_text(slide.shapes[17], "85.71%")
    set_shape_text(slide.shapes[18], "21.3M")
    set_shape_text(slide.shapes[20], "MobileNetV3")
    set_shape_text(slide.shapes[21], "49.08%")
    set_shape_text(slide.shapes[22], "1.5M")
    set_shape_text(
        slide.shapes[23],
        "ResNet18 показал лучший результат по Accuracy среди протестированных моделей.",
    )

    old_chart = slide.shapes[25]
    chart_left = old_chart.left
    chart_top = old_chart.top
    chart_width = old_chart.width
    chart_height = old_chart.height
    remove_shape(old_chart)
    slide.shapes.add_picture(str(CHART_PATH), chart_left, chart_top, chart_width, chart_height)

    # Slide 8
    slide = prs.slides[7]
    set_shape_text(slide.shapes[2], "ResNet18")
    set_shape_text(slide.shapes[5], "89.7%")
    set_shape_text(slide.shapes[8], "88.2%")
    set_shape_text(slide.shapes[11], "11.2M")
    set_shape_text(
        slide.shapes[15],
        "Финальная модель по обновлённому сравнению, 20 эпох",
    )
    slide.shapes[15].width = 7000000
    style_shape_text(slide.shapes[15], font_size=13)

    # Slide 9
    slide = prs.slides[8]
    set_shape_text(slide.shapes[9], "Загрузка весов финальной модели")

    # Slide 10
    slide = prs.slides[9]
    set_shape_text(slide.shapes[10], "Проведено обновлённое сравнение 3 архитектур (Этап 5, 20 эпох)")
    set_shape_text(slide.shapes[12], "Финальная модель ResNet18: accuracy 89.7%")
    set_shape_text(slide.shapes[14], "Подготовлен ноутбук для инференса с загрузкой своих изображений")
    set_shape_text(slide.shapes[15], "Текущие ограничения: редкие классы и визуально похожие дефекты.")
    set_shape_text(slide.shapes[17], "Дальнейшая работа")
    set_shape_text(slide.shapes[25], "Тема №9")
    set_shape_text(slide.shapes[26], "Классификация дефектов на поверхности стали")

    lower_card_bg = slide.shapes[24]
    theme_title = slide.shapes[25]
    theme_subtitle = slide.shapes[26]

    old_right_shapes = [18, 19, 20, 21, 22, 23]
    for idx in sorted(old_right_shapes, reverse=True):
        remove_shape(slide.shapes[idx])

    add_textbox(slide, 5582580, 1588554, 2733675, 176212, "Балансировка классов", font_size=12, bold=True, color=(241, 245, 249))
    add_textbox(slide, 5582580, 1785000, 2550000, 360000, "Улучшить качество на редких дефектах.", font_size=11.5, color=(148, 163, 184))
    add_textbox(slide, 5582580, 2120000, 2733675, 176212, "Аугментации", font_size=12, bold=True, color=(241, 245, 249))
    add_textbox(slide, 5582580, 2315000, 2550000, 420000, "Сделать модель устойчивее к вариативности изображений.", font_size=11.5, color=(148, 163, 184))
    add_textbox(slide, 5582580, 2725000, 2733675, 176212, "Новые архитектуры", font_size=12, bold=True, color=(241, 245, 249))
    add_textbox(slide, 5582580, 2920000, 2550000, 420000, "Сравнить более сильные CNN и доработать inference-пайплайн.", font_size=11.5, color=(148, 163, 184))

    lower_card_bg.top = 3400000
    theme_title.top = 3520000
    theme_subtitle.top = 3890000

    if TEMP_OUTPUT_PPTX.exists():
        TEMP_OUTPUT_PPTX.unlink()
    prs.save(str(TEMP_OUTPUT_PPTX))
    if OUTPUT_PPTX.exists():
        OUTPUT_PPTX.unlink()
    shutil.move(str(TEMP_OUTPUT_PPTX), str(OUTPUT_PPTX))
    print(OUTPUT_PPTX)
    print(CHART_PATH)


if __name__ == "__main__":
    main()
