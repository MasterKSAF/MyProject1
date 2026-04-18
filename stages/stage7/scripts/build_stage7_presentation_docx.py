from pathlib import Path
import re

from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[3]
STAGE_DIR = ROOT / "stages" / "stage7"
SOURCE_PATH = STAGE_DIR / "docs" / "stage7_presentation_slides_ru.md"
OUTPUT_PATH = STAGE_DIR / "docs" / "stage7_presentation_slides.docx"


def normalize_line(line: str) -> str:
    line = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1: \2", line)
    line = line.replace("`", "")
    return line


def main() -> None:
    text = SOURCE_PATH.read_text(encoding="utf-8")
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    for raw_line in text.splitlines():
        line = raw_line.rstrip()
        if line.startswith("# "):
            document.add_heading(normalize_line(line[2:].strip()), level=1)
            continue
        if line.startswith("## "):
            document.add_heading(normalize_line(line[3:].strip()), level=2)
            continue
        if not line:
            document.add_paragraph("")
            continue
        if line.startswith("- "):
            document.add_paragraph(normalize_line(line[2:].strip()), style="List Bullet")
            continue
        document.add_paragraph(normalize_line(line.strip()))

    document.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    main()
